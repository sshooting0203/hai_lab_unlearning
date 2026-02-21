from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from datetime import datetime
import os 

class ExcelLogger:
    def __init__(self, xlsx_path: str, sheet_name: str = "log"):
        self.xlsx_path = xlsx_path
        self.sheet_name = sheet_name

        if os.path.exists(self.xlsx_path):
            self.wb = load_workbook(self.xlsx_path)
            if self.sheet_name not in self.wb.sheetnames:
                self.ws = self.wb.create_sheet(self.sheet_name)
                self._write_header()
            else:
                self.ws = self.wb[self.sheet_name]
        else:
            os.makedirs(os.path.dirname(self.xlsx_path), exist_ok=True)
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = self.sheet_name
            self._write_header()
            self.wb.save(self.xlsx_path)

    def _write_header(self):
        headers = [
            "seed",
            "unlearn_dataset",
            "auroc",
            "w_mean_train",
            "w_mean_test",
            "w_mean_ood",
            "threshold",
            "x0",
            "acc",
            "best_acc_global",
        ]
        self.ws.append(headers)
        self.wb.save(self.xlsx_path)

    def log_row(self, row: dict):
        # header 순서대로 append
        ordered = [
            row.get("seed"),
            row.get("unlearn_dataset"),
            row.get("auroc"),
            row.get("w_mean_train"),
            row.get("w_mean_test"),
            row.get("w_mean_ood"),
            row.get("threshold"),
            row.get("x0"),
            row.get("acc"),
            row.get("best_acc_global"),
        ]
        self.ws.append(ordered)
        self.wb.save(self.xlsx_path)
